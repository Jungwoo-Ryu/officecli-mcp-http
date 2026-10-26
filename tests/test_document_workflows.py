"""Real document workflows; optionally target the deployed Docker container.

OFFICECLI_WORKFLOW_URL=http://localhost:18080 runs these against Docker.
OFFICECLI_ARTIFACT_DIR saves template inputs and downloaded outputs for inspection.
"""
from __future__ import annotations

import json
import os
from contextlib import asynccontextmanager
from copy import copy
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from mcp import ClientSession
from mcp.client.streamable_http import streamable_http_client
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from officecli_mcp.core import Settings


@pytest.fixture
def workflow_server(request):
    external = os.getenv("OFFICECLI_WORKFLOW_URL")
    if external:
        settings = Settings()
        return external.removesuffix("/mcp").rstrip("/"), settings
    return request.getfixturevalue("running")


@asynccontextmanager
async def workflow(server):
    base, settings = server
    auth = {"Authorization": f"Bearer {settings.token}", "X-User-Id": f"workflow-{uuid4().hex[:12]}"}
    async with (
        httpx.AsyncClient(base_url=base, headers=auth, timeout=60) as http,
        httpx.AsyncClient(headers=auth, timeout=60) as transport,
        streamable_http_client(f"{base}/mcp", http_client=transport) as (read, write, _),
        ClientSession(read, write) as session,
    ):
        await session.initialize()

        async def call(name, **args):
            result = await session.call_tool(name, args)
            assert not result.isError, result
            return result.structuredContent or json.loads(result.content[0].text)

        yield http, call


def artifact(name, data):
    if directory := os.getenv("OFFICECLI_ARTIFACT_DIR"):
        target = Path(directory) / name
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)


def word_template():
    doc = Document()
    doc.core_properties.title = "회사 표준 견적서"
    doc.sections[0].top_margin = Cm(2.5)
    doc.sections[0].left_margin = Cm(2.2)
    doc.sections[0].header.paragraphs[0].text = "ACME | 회사 표준 양식"
    doc.sections[0].footer.paragraphs[0].text = "CONFIDENTIAL | 승인 없이 배포 금지"
    doc.add_heading("표준 견적서", level=1)
    p = doc.add_paragraph()
    p.add_run("수신: ").bold = True
    run = p.add_run("{{고객명}}")
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("1F4E78")
    run.italic = True
    p.add_run(" 귀하").underline = True
    doc.add_paragraph("변경하면 안 되는 계약 조건", style="Quote")
    table = doc.add_table(rows=2, cols=2, style="Light Shading Accent 1")
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "내용"
    table.cell(1, 0).text = "프로젝트"
    table.cell(1, 1).text = "{{프로젝트명}}"
    table.cell(1, 1).paragraphs[0].runs[0].bold = True
    # A bookmark and its ID must survive edits elsewhere in the document.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "41")
    start.set(qn("w:name"), "unchanged_clause")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "41")
    doc.paragraphs[2]._p.insert(0, start)
    doc.paragraphs[2]._p.append(end)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def excel_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "견적서"
    config = wb.create_sheet("설정")
    config["A1"] = "부가세율"
    config["B1"] = 0.1
    config["B1"].number_format = "0%"
    ws.merge_cells("A1:D1")
    ws["A1"] = "회사 표준 견적서"
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A3"] = "고객명"
    ws["B3"] = "{{고객명}}"
    for row in [["품목", "수량", "단가", "금액"], ["개발 용역", 2, 100000, "=B5*C5"]]:
        ws.append(row)  # rows 4 and 5
    ws["D6"] = "=D5*'설정'!B1"
    ws["D7"] = "=SUM(D5:D6)"
    ws["A9"] = "이 문구는 그대로 유지"
    ws["B5"].font = Font(name="Arial", color="0000FF")
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in ws.iter_rows(min_row=5, max_row=7, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '#,##0"원"'
    ws.column_dimensions["A"].width = 27
    ws.column_dimensions["B"].width = 16
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B5"
    ws.print_area = "A1:D9"
    ws.page_setup.orientation = "landscape"
    validation = DataValidation(type="whole", operator="between", formula1=1, formula2=100)
    ws.add_data_validation(validation)
    validation.add(ws["B5"])
    chart = BarChart()
    chart.title = "항목별 금액"
    chart.add_data(Reference(ws, min_col=4, min_row=4, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=5))
    ws.add_chart(chart, "F2")
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


@pytest.mark.asyncio
async def test_new_word_document_with_table(workflow_server):
    async with workflow(workflow_server) as (http, call):
        await call("office_help", format="docx", topic="table")
        result = await call("office_create", filename="신규-제안서.docx", operations=[
            {"action": "add", "path": "/body", "element_type": "paragraph", "props": {"text": "2026년 개발 제안서", "style": "Heading1"}},
            {"action": "add", "path": "/body", "element_type": "paragraph", "props": {"text": "고객 요구사항에 따른 신규 문서입니다.", "bold": True}},
            {"action": "add", "path": "/body", "element_type": "table", "props": {"rows": 2, "cols": 2, "data": "항목,금액;개발비,1500000"}},
        ])
        response = await http.get(result["download_url"])
        assert response.status_code == 200
        data = response.content
        doc = Document(BytesIO(data))
        title = next(p for p in doc.paragraphs if p.text == "2026년 개발 제안서")
        assert title.style.style_id == "Heading1"
        assert any(p.text == "고객 요구사항에 따른 신규 문서입니다." for p in doc.paragraphs)
        assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["항목", "금액"], ["개발비", "1500000"]]
        artifact("신규-제안서.docx", data)


@pytest.mark.asyncio
async def test_new_excel_multisheet_and_formulas(workflow_server):
    async with workflow(workflow_server) as (http, call):
        await call("office_help", format="xlsx", topic="sheet")
        result = await call("office_create", filename="신규-예산.xlsx", operations=[
            {"action": "set", "path": "/Sheet1/A1", "props": {"value": "수량", "bold": True, "fill": "4472C4"}},
            {"action": "set", "path": "/Sheet1/A2", "props": {"value": 3}},
            {"action": "set", "path": "/Sheet1/B2", "props": {"value": 50000}},
            {"action": "set", "path": "/Sheet1/C2", "props": {"formula": "A2*B2", "numberformat": "#,##0"}},
            {"action": "add", "path": "/", "element_type": "sheet", "props": {"name": "요약"}},
            {"action": "set", "path": "/요약/A1", "props": {"value": "합계"}},
            {"action": "set", "path": "/요약/B1", "props": {"formula": "Sheet1!C2"}},
        ])
        data = (await http.get(result["download_url"])).content
        wb = load_workbook(BytesIO(data), data_only=False)
        assert wb.sheetnames == ["Sheet1", "요약"]
        assert wb["Sheet1"]["C2"].value == "=A2*B2"
        assert wb["Sheet1"]["C2"].number_format == "#,##0"
        assert wb["Sheet1"]["A1"].font.bold is True
        assert wb["Sheet1"]["A1"].fill.fgColor.rgb.endswith("4472C4")
        computed = await call("office_get", filename=result["filename"], path="/요약/B1")
        assert computed["result"]["data"]["results"][0]["format"]["computedValue"] == "150000"
        artifact("신규-예산.xlsx", data)


@pytest.mark.asyncio
async def test_word_template_preserves_layout_runs_and_table(workflow_server):
    source = word_template()
    artifact("템플릿-견적서.docx", source)
    async with workflow(workflow_server) as (http, call):
        uploaded = await http.post("/api/files", files={"file": ("템플릿-견적서.docx", source)})
        assert uploaded.status_code == 201, uploaded.text
        nodes = await call("office_get", filename=uploaded.json()["filename"], path="/body", depth=5)
        assert "{{고객명}}" in json.dumps(nodes, ensure_ascii=False)
        result = await call("office_edit", filename=uploaded.json()["filename"], output_filename="작성완료-견적서.docx", operations=[
            {"action": "set", "path": "/body/p[2]/r[2]", "props": {"text": "테스트 주식회사"}},
            {"action": "set", "path": "/body/tbl[1]/tr[2]/tc[2]/p[1]/r[1]", "props": {"text": "웹앱 구축 프로젝트"}},
        ])
        data = (await http.get(result["download_url"])).content
        doc, before = Document(BytesIO(data)), Document(BytesIO(source))
        assert doc.paragraphs[1].text == "수신: 테스트 주식회사 귀하"
        assert doc.paragraphs[1].runs[0].bold is True
        assert doc.paragraphs[1].runs[1].italic is True
        assert doc.paragraphs[1].runs[1].font.size == Pt(13)
        assert str(doc.paragraphs[1].runs[1].font.color.rgb) == "1F4E78"
        assert doc.paragraphs[1].runs[2].underline is True
        assert doc.paragraphs[2].text == before.paragraphs[2].text
        assert doc.paragraphs[2].style.style_id == before.paragraphs[2].style.style_id
        assert doc.tables[0].cell(1, 1).text == "웹앱 구축 프로젝트"
        assert doc.tables[0].cell(1, 1).paragraphs[0].runs[0].bold is True
        assert doc.tables[0].style.style_id == before.tables[0].style.style_id
        assert doc.sections[0].header.paragraphs[0].text == before.sections[0].header.paragraphs[0].text
        assert doc.sections[0].footer.paragraphs[0].text == before.sections[0].footer.paragraphs[0].text
        assert doc.sections[0].top_margin == before.sections[0].top_margin
        assert doc.sections[0].left_margin == before.sections[0].left_margin
        assert doc.core_properties.title == before.core_properties.title
        with ZipFile(BytesIO(data)) as archive:
            assert b"unchanged_clause" in archive.read("word/document.xml")
        assert (await http.get("/api/files/템플릿-견적서.docx")).content == source
        artifact("작성완료-견적서.docx", data)
        # A follow-up edits the returned copy and preserves the first result.
        followup = await call("office_edit", filename=result["filename"], operations=[
            {"action": "set", "path": "/body/p[2]/r[2]", "props": {"text": "후속 수정 고객사"}},
        ])
        followup_data = (await http.get(followup["download_url"])).content
        assert "후속 수정 고객사" in Document(BytesIO(followup_data)).paragraphs[1].text
        assert (await http.get("/api/files/작성완료-견적서.docx")).content == data
        artifact("후속수정-견적서.docx", followup_data)


@pytest.mark.asyncio
async def test_excel_template_preserves_styles_chart_and_rules(workflow_server):
    source = excel_template()
    artifact("템플릿-예산.xlsx", source)
    async with workflow(workflow_server) as (http, call):
        uploaded = await http.put("/api/files/템플릿-예산.xlsx", content=source)
        assert uploaded.status_code == 201, uploaded.text
        await call("office_get", filename=uploaded.json()["filename"], path="/견적서", depth=2)
        result = await call("office_edit", filename=uploaded.json()["filename"], output_filename="작성완료-예산.xlsx", operations=[
            {"action": "set", "path": "/견적서/B3", "props": {"value": "테스트 주식회사"}},
            {"action": "set", "path": "/견적서/B5", "props": {"value": 4}},
        ])
        data = (await http.get(result["download_url"])).content
        wb, before = load_workbook(BytesIO(data)), load_workbook(BytesIO(source))
        ws, old = wb["견적서"], before["견적서"]
        assert wb.sheetnames == before.sheetnames
        assert ws["B3"].value == "테스트 주식회사"
        assert ws["B5"].value == 4
        assert ws["A9"].value == old["A9"].value
        assert wb["설정"]["B1"].value == before["설정"]["B1"].value
        assert ws["D5"].value == old["D5"].value
        assert ws["D6"].value == old["D6"].value
        assert ws["D7"].value == old["D7"].value
        for address in ["A1", "A4", "B5", "C5", "D5"]:
            assert copy(ws[address].font) == copy(old[address].font)
            assert copy(ws[address].fill) == copy(old[address].fill)
            assert ws[address].number_format == old[address].number_format
            assert copy(ws[address].alignment) == copy(old[address].alignment)
        assert str(ws.merged_cells) == str(old.merged_cells)
        assert ws.column_dimensions["A"].width == old.column_dimensions["A"].width
        assert ws.row_dimensions[1].height == old.row_dimensions[1].height
        assert ws.freeze_panes == old.freeze_panes
        assert str(ws.print_area) == str(old.print_area)
        assert ws.page_setup.orientation == old.page_setup.orientation
        assert ws.data_validations.to_tree().attrib == old.data_validations.to_tree().attrib
        assert len(ws.data_validations.dataValidation) == 1
        rule, old_rule = ws.data_validations.dataValidation[0], old.data_validations.dataValidation[0]
        assert (rule.type, rule.operator, rule.formula1, rule.formula2, str(rule.sqref)) == (
            old_rule.type, old_rule.operator, old_rule.formula1, old_rule.formula2, str(old_rule.sqref))
        assert len(ws._charts) == len(old._charts) == 1
        assert ws._charts[0].series[0].val.numRef.f == old._charts[0].series[0].val.numRef.f
        assert (ws._charts[0].anchor._from.col, ws._charts[0].anchor._from.row) == (
            old._charts[0].anchor._from.col, old._charts[0].anchor._from.row)
        computed = await call("office_get", filename=result["filename"], path="/견적서/D7")
        assert computed["result"]["data"]["results"][0]["format"]["computedValue"] == "440000"
        assert (await http.get("/api/files/템플릿-예산.xlsx")).content == source
        artifact("작성완료-예산.xlsx", data)


@pytest.mark.asyncio
@pytest.mark.parametrize("filename", ["native-template.dotx", "native-template.xltx"])
async def test_native_template_extensions_are_explicitly_unsupported(workflow_server, filename):
    async with workflow(workflow_server) as (http, _):
        response = await http.put(f"/api/files/{filename}", content=word_template())
        assert response.status_code == 400
        assert "docx/.xlsx" in response.json()["error"]
