from __future__ import annotations

import json
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import pytest
from docx import Document
from mcp import ClientSession
from mcp.client.streamable_http import streamablehttp_client
from openpyxl import load_workbook

from officecli_mcp.core import Store


def headers(settings, user="alice"):
    return {"Authorization": f"Bearer {settings.token}", "X-User-Id": user}


def docx():
    document = Document()
    document.add_paragraph("첨부 원본 문장")
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_mcp_attachment_and_generation(running):
    base, settings = running
    auth = headers(settings)
    async with httpx.AsyncClient(base_url=base, timeout=30) as http:
        uploaded = await http.post("/api/files", headers=auth,
                                   files={"file": ("첨부.docx", docx())})
        assert uploaded.status_code == 201, uploaded.text
        async with (
            streamablehttp_client(f"{base}/mcp", headers=auth) as (read, write, _),
            ClientSession(read, write) as session,
        ):
            await session.initialize()
            assert "office_edit" in {t.name for t in (await session.list_tools()).tools}

            async def call(name, **args):
                result = await session.call_tool(name, args)
                assert not result.isError, result
                return result.structuredContent or json.loads(result.content[0].text)

            await call("office_help", format="docx", topic="paragraph")
            initial = await call("office_read", filename="첨부.docx")
            assert "첨부 원본 문장" in json.dumps(initial, ensure_ascii=False)
            edited = await call("office_edit", filename="첨부.docx", output_filename="수정본.docx",
                                operations=[{"action": "set", "path": "/body/p[1]", "props": {"text": "수정된 문장"}}])
            response = await http.get(edited["download_url"])
            assert response.status_code == 200
            assert "filename*=UTF-8''" in response.headers["content-disposition"]
            parsed = Document(BytesIO(response.content))
            assert parsed.paragraphs[0].text == "수정된 문장"
            original = await http.get("/api/files/첨부.docx", headers=auth)
            assert Document(BytesIO(original.content)).paragraphs[0].text == "첨부 원본 문장"
            assert (await http.get("/api/files/첨부.docx", headers=headers(settings, "bob"))).status_code == 404
            await call("office_create", filename="생성.docx", operations=[{"action": "add", "path": "/body", "element_type": "paragraph", "props": {"text": "새 보고서"}}])
            book = await call("office_create", filename="예산.xlsx", operations=[
                {"action": "set", "path": "/Sheet1/A1", "props": {"value": 10}},
                {"action": "set", "path": "/Sheet1/B1", "props": {"formula": "A1*2"}},
            ])
            book_bytes = (await http.get(book["download_url"])).content
            assert load_workbook(BytesIO(book_bytes), data_only=True)["Sheet1"]["B1"].value == 20
            # XLSX attachments use the same upload and edit flow.
            assert (await http.put("/api/files/uploaded.xlsx", headers=auth, content=book_bytes)).status_code == 201
            updated = await call("office_edit", filename="uploaded.xlsx", output_filename="예산-수정.xlsx",
                                 operations=[{"action": "set", "path": "/Sheet1/A1", "props": {"value": 30}}])
            changed = (await http.get(updated["download_url"])).content
            # OfficeCLI invalidates the old formula cache on input changes.
            # Its evaluator supplies computedValue; Excel refreshes caches on open.
            computed = await call("office_get", filename="예산-수정.xlsx", path="/Sheet1/B1")
            assert computed["result"]["data"]["results"][0]["format"]["computedValue"] == "60"
            assert load_workbook(BytesIO(changed), data_only=True)["Sheet1"]["B1"].value in (None, 60)
            assert load_workbook(BytesIO(changed), data_only=False)["Sheet1"]["B1"].value == "=A1*2"
            # A failed mutation publishes no partial output and leaves source intact.
            failure = await session.call_tool("office_edit", {"filename": "첨부.docx", "output_filename": "failed.docx", "operations": [
                {"action": "set", "path": "/body/p[1]", "props": {"text": "should not persist"}},
                {"action": "set", "path": "/body/p[999]", "props": {"text": "missing"}},
            ]})
            assert failure.isError
            assert (await http.get("/api/files/failed.docx", headers=auth)).status_code == 404
            assert (await http.get("/api/files/첨부.docx", headers=auth)).content == original.content
        # A separate MCP caller must not see Alice's attachments.
        async with (
            streamablehttp_client(f"{base}/mcp", headers=headers(settings, "bob")) as (r, w, _),
            ClientSession(r, w) as session,
        ):
            await session.initialize()
            result = await session.call_tool("office_read", {"filename": "첨부.docx"})
            assert result.isError


def test_http_guards_and_upload_limits(running):
    base, settings = running
    auth = headers(settings)
    with httpx.Client(base_url=base) as http:
        assert http.get("/health").status_code == 200
        assert http.post("/mcp", json={}).status_code == 401
        assert http.put("/api/files/a.docx", content=docx()).status_code == 401
        assert http.get("/api/files/a.docx", headers={"Authorization": f"Bearer {settings.token}"}).status_code == 400
        assert http.put("/api/files/a.docx", headers=auth, content=b"invalid").status_code == 400
        assert http.put("/api/files/a.docx", headers=auth, content=docx()).status_code == 201
        assert http.put("/api/files/a.docx", headers=auth, content=docx()).status_code == 409
        assert http.put("/api/files/a.docx", headers=auth, content=b"x" * (1024 * 1024 + 1)).status_code == 413
        assert http.put("/api/files/b.docx", headers=auth, content=iter([b"x" * 700000, b"y" * 700000])).status_code == 413
        token = Store(settings).metadata("alice", "a.docx")["download_url"]
        assert http.get(token + "x").status_code == 403


@pytest.mark.parametrize("filename", ["../a.docx", "/a.docx", "a/b.xlsx", "x\\a.docx", ".hidden.docx", "a.exe"])
def test_path_sandbox(settings, filename):
    with pytest.raises(ValueError):
        Store(settings).path("alice", filename)


def test_archive_bomb_and_bad_user(settings):
    store = Store(settings)
    with pytest.raises(ValueError):
        store.path("..", "a.docx")
    data = BytesIO()
    with ZipFile(data, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"a" * (11 * 1024 * 1024))
    with pytest.raises(ValueError, match="expansion"):
        store.validate(data.getvalue(), "a.docx")


def test_operations_disallow_external_files():
    from officecli_mcp.server import Operation
    with pytest.raises(ValueError):
        Operation(action="add", path="/body", element_type="picture", props={"src": "/etc/passwd"})
    with pytest.raises(ValueError):
        Operation(action="set", path="/body", props={"src": "/etc/passwd"})
